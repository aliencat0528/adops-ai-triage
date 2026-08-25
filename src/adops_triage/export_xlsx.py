# -*- coding: utf-8 -*-
"""
把 tickets.csv 匯出成給人看的多工作表 xlsx。

工作表：
  1. 需求單主檔   中文欄名、凍結標題列、自動篩選、條件式色階
  2. 欄位字典     英文鍵 ↔ 中文欄名 ↔ 群組 ↔ 型別 ↔ 來源 ↔ 設計理由
  3. 下拉選單清單 各列舉欄位的合法值（可直接做資料驗證）
  4. 統計摘要     以公式計算，改動主檔資料會自動重算

用法：PYTHONPATH=src python -m adops_triage.export_xlsx
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

CSV = Path("data/raw/tickets.csv")
OUT = Path("data/raw/廣告支援需求單.xlsx")

FONT = "Arial"

# (英文鍵, 中文欄名, 群組, 型別, 來源/計算, 設計理由)
FIELDS = [
    ("ticket_id", "工單編號", "A 工單識別", "字串", "系統自動", "REQ-YYYYMM-NNNN"),
    ("submitted_at", "提單時間", "A 工單識別", "日期時間", "系統自動", "季節性與時段分析的基礎"),
    ("org_context", "組織情境", "A 工單識別", "列舉", "taxonomy.yaml", "讓三種公司型態可切片比較"),
    ("requester_name", "提單人", "A 工單識別", "字串", "登入身分", "追溯與教育訓練對象"),
    (
        "requester_role",
        "提單人角色",
        "A 工單識別",
        "列舉",
        "組織資料",
        "角色是提單品質最強的預測因子",
    ),
    ("client_brand", "客戶／品牌", "B 客戶脈絡", "字串", "客戶主檔", "找出反覆出問題的客戶"),
    ("industry", "產業別", "B 客戶脈絡", "列舉", "客戶主檔", "不同產業的問題結構不同"),
    ("sla_tier", "服務等級SLA", "B 客戶脈絡", "列舉", "合約資料", "決定首次回應時間目標"),
    (
        "website_platform",
        "官網／建站平台",
        "B 客戶脈絡",
        "列舉",
        "客戶主檔",
        "建站平台直接決定根因分布",
    ),
    (
        "issue_category",
        "問題大類",
        "C 問題描述",
        "列舉",
        "taxonomy.yaml",
        "分診的第一層，決定動態必填欄位",
    ),
    (
        "issue_subcategory",
        "問題細類",
        "C 問題描述",
        "列舉",
        "taxonomy.yaml",
        "對應知識庫文章與診斷探針",
    ),
    ("title", "工單標題", "C 問題描述", "字串", "提單人填寫", "分診模型的輸入之一"),
    ("description", "問題描述", "C 問題描述", "長文字", "提單人填寫", "完整度評分與分診的主要輸入"),
    ("platform", "廣告平台", "C 問題描述", "列舉", "taxonomy.yaml", "決定要呼叫哪一組平台探針"),
    (
        "affected_asset_id",
        "受影響資產ID",
        "C 問題描述",
        "字串",
        "提單人填寫",
        "缺了診斷無法進行，權重最高的必填欄位",
    ),
    (
        "tracking_method",
        "追蹤方式",
        "D 技術上下文",
        "列舉",
        "提單人填寫",
        "Client/Server/Hybrid 決定除錯路徑",
    ),
    ("environment", "環境", "D 技術上下文", "列舉", "提單人填寫", "正式站與測試站的處理優先級不同"),
    (
        "reproducible",
        "可重現性",
        "D 技術上下文",
        "列舉",
        "提單人填寫",
        "無法重現會直接增加一次釐清",
    ),
    ("occurred_from", "異常起始日", "D 技術上下文", "日期", "提單人填寫", "決定探針的查詢區間"),
    ("severity", "嚴重度", "E 影響量化", "列舉", "分診判定", "S1-S4，依投放中斷程度"),
    ("priority", "優先級", "E 影響量化", "列舉", "嚴重度 × SLA", "P0-P3，排隊順序"),
    (
        "impacted_spend_daily_ntd",
        "受影響日花費NTD",
        "E 影響量化",
        "整數",
        "廣告後台",
        "換算商業影響金額",
    ),
    (
        "impacted_campaign_count",
        "受影響活動數",
        "E 影響量化",
        "整數",
        "廣告後台",
        "影響面計算的乘數",
    ),
    (
        "metric_name",
        "關鍵指標名稱",
        "E 影響量化",
        "字串",
        "依問題大類",
        "讓不同類問題可用同一套欄位量化",
    ),
    (
        "metric_before",
        "異常前數值",
        "E 影響量化",
        "數值",
        "提單人填寫",
        "沒有前後值就無法判斷是驟降還是自然衰減",
    ),
    ("metric_after", "異常後數值", "E 影響量化", "數值", "提單人填寫", "同上"),
    ("delta_pct", "變化幅度%", "E 影響量化", "數值", "自動計算", "嚴重度判定的輸入"),
    ("business_impact", "商業影響", "E 影響量化", "列舉", "提單人填寫", "區分技術問題與商業後果"),
    ("attachment_count", "附件數", "F 提單品質", "整數", "系統統計", "完整度的代理指標之一"),
    (
        "info_completeness_score",
        "資訊完整度分數",
        "F 提單品質",
        "0-100",
        "Intake Agent 計算",
        "★ 核心欄位，與釐清次數相關 -0.727",
    ),
    ("missing_fields", "缺漏資訊欄位", "F 提單品質", "字串", "Intake Agent 計算", "追問內容的依據"),
    ("first_response_min", "首次回應時間(分)", "F 提單品質", "整數", "系統計算", "SLA 達成率"),
    (
        "clarification_rounds",
        "來回釐清次數",
        "F 提單品質",
        "整數",
        "系統統計",
        "★ 最大成本放大器，每次 +22.2% 處理時數",
    ),
    (
        "is_duplicate",
        "是否重複工單",
        "F 提單品質",
        "列舉",
        "分診判定",
        "相似度 >0.88 且同客戶 7 天內",
    ),
    ("duplicate_of", "重複來源工單", "F 提單品質", "字串", "分診判定", "合併後的主單號"),
    ("assigned_team", "承接團隊", "G 處理與根因", "列舉", "分診路由", "路由錯誤率的評測依據"),
    ("assignee", "處理工程師", "G 處理與根因", "字串", "指派", "負載與專長分析"),
    (
        "root_cause_category",
        "根因大類",
        "G 處理與根因",
        "列舉",
        "taxonomy.yaml",
        "★ 判斷可預防性的依據",
    ),
    (
        "root_cause_code",
        "根因代碼",
        "G 處理與根因",
        "字串",
        "taxonomy.yaml",
        "程式用鍵值，對應偵測器",
    ),
    (
        "resolution_type",
        "處理方式",
        "G 處理與根因",
        "列舉",
        "taxonomy.yaml",
        "教育說明佔比 = 知識庫可解的比例",
    ),
    ("fix_owner", "修復責任方", "G 處理與根因", "列舉", "自動推導", "區分內部問題與客戶端問題"),
    ("status", "狀態", "G 處理與根因", "列舉", "工單狀態機", "結案率與積壓分析"),
    ("resolved_at", "結案時間", "G 處理與根因", "日期時間", "系統自動", "TTR 計算"),
    ("resolution_hours", "總處理時數", "G 處理與根因", "數值", "系統計算", "含等待時間的完整耗時"),
    (
        "eng_effort_hours",
        "工程實際投入工時",
        "G 處理與根因",
        "數值",
        "工程回報",
        "★ 北極星指標的分子",
    ),
    ("reopened_count", "重啟次數", "G 處理與根因", "整數", "系統統計", "真實解決 vs 假結案的判別"),
    (
        "recurrence_90d",
        "90天內復發",
        "G 處理與根因",
        "列舉",
        "系統計算",
        "同客戶同類問題，指向結構性根因",
    ),
    ("csat", "滿意度CSAT", "G 處理與根因", "1-5", "結案問卷", "與釐清次數單調負相關"),
    (
        "ai_automation_level",
        "AI自動化分級",
        "H AI 標註",
        "列舉",
        "規則標註",
        "★ L0-L4，決定投資順序",
    ),
    ("ai_preventable", "可否主動偵測預防", "H AI 標註", "列舉", "根因推導", "48.9% 標註為可預防"),
    ("detection_signal", "建議偵測訊號", "H AI 標註", "字串", "根因推導", "直接對應到六個偵測器"),
    ("kb_article_id", "對應知識庫文章", "H AI 標註", "字串", "細類對映", "檢索層的錨點"),
    (
        "platform_event_id",
        "平台事件錨點",
        "H AI 標註",
        "字串",
        "生成器標註",
        "標示屬於哪個 2026 真實平台事件",
    ),
]

ZH = {en: zh for en, zh, *_ in FIELDS}

HEAD_FILL = PatternFill("solid", fgColor="1F3243")
HEAD_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT, size=10)
TITLE_FONT = Font(name=FONT, size=13, bold=True, color="1F3243")
NOTE_FONT = Font(name=FONT, size=9, italic=True, color="666666")
KEY_FILL = PatternFill("solid", fgColor="FBE7DA")
THIN = Side(style="thin", color="D5DCE3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, ncols: int, row: int = 1) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 30


def main() -> None:
    df = pd.read_csv(CSV)
    wb = Workbook()

    # ---------------------------------------------------- 1. 主檔
    ws = wb.active
    ws.title = "需求單主檔"
    cols = [f[0] for f in FIELDS if f[0] in df.columns]
    ws.append([ZH[c] for c in cols])
    for _, r in df[cols].iterrows():
        ws.append(list(r.values))
    style_header(ws, len(cols))

    widths = {
        "工單標題": 34,
        "問題描述": 62,
        "缺漏資訊欄位": 26,
        "建議偵測訊號": 34,
        "受影響資產ID": 24,
        "問題細類": 30,
        "根因大類": 26,
        "AI自動化分級": 18,
        "客戶／品牌": 16,
        "官網／建站平台": 18,
        "服務等級SLA": 18,
        "組織情境": 22,
        "處理方式": 20,
        "承接團隊": 18,
        "商業影響": 18,
        "關鍵指標名稱": 16,
    }
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(ZH[c], 13)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(cols)):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"

    # 完整度分數色階（紅→綠），讓提單品質一眼可見
    comp_idx = cols.index("info_completeness_score") + 1
    comp_col = get_column_letter(comp_idx)
    ws.conditional_formatting.add(
        f"{comp_col}2:{comp_col}{ws.max_row}",
        ColorScaleRule(
            start_type="num",
            start_value=0,
            start_color="E8A79B",
            mid_type="num",
            mid_value=56,
            mid_color="F5E2C8",
            end_type="num",
            end_value=100,
            end_color="A9CDB6",
        ),
    )
    clar_col = get_column_letter(cols.index("clarification_rounds") + 1)
    ws.conditional_formatting.add(
        f"{clar_col}2:{clar_col}{ws.max_row}",
        ColorScaleRule(
            start_type="num",
            start_value=0,
            start_color="A9CDB6",
            end_type="num",
            end_value=6,
            end_color="E8A79B",
        ),
    )

    n_rows = ws.max_row

    # ---------------------------------------------------- 2. 欄位字典
    d = wb.create_sheet("欄位字典")
    d.append(["英文鍵（程式用）", "中文欄名（人看）", "群組", "型別", "來源／計算方式", "設計理由"])
    for en, zh, grp, typ, src, why in FIELDS:
        d.append([en, zh, grp, typ, src, why])
    style_header(d, 6)
    for w, c in zip([28, 22, 16, 12, 22, 52], "ABCDEF"):
        d.column_dimensions[c].width = w
    for row in d.iter_rows(min_row=2, max_row=d.max_row, max_col=6):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
        if str(row[5].value).startswith("★"):
            for cell in row:
                cell.fill = KEY_FILL
    d.freeze_panes = "A2"
    d.append([])
    d.append(
        [
            "★ 標記者為本專案的核心欄位——一般工單系統沒有這三組欄位，"
            "因此答不出「為什麼花這麼久」。",
            "",
            "",
            "",
            "",
            "",
        ]
    )
    d.cell(row=d.max_row, column=1).font = NOTE_FONT

    # ---------------------------------------------------- 3. 下拉選單清單
    v = wb.create_sheet("下拉選單清單")
    enum_cols = [
        "org_context",
        "requester_role",
        "industry",
        "sla_tier",
        "website_platform",
        "issue_category",
        "platform",
        "tracking_method",
        "environment",
        "reproducible",
        "severity",
        "priority",
        "business_impact",
        "assigned_team",
        "root_cause_category",
        "resolution_type",
        "fix_owner",
        "status",
        "ai_automation_level",
        "ai_preventable",
    ]
    enum_cols = [c for c in enum_cols if c in df.columns]
    v.append([ZH[c] for c in enum_cols])
    style_header(v, len(enum_cols))
    vals = {c: sorted(df[c].dropna().astype(str).unique().tolist()) for c in enum_cols}
    for i in range(max(len(x) for x in vals.values())):
        v.append([vals[c][i] if i < len(vals[c]) else "" for c in enum_cols])
    for i, c in enumerate(enum_cols, start=1):
        v.column_dimensions[get_column_letter(i)].width = 24
    for row in v.iter_rows(min_row=2, max_row=v.max_row, max_col=len(enum_cols)):
        for cell in row:
            cell.font = BODY_FONT
            cell.border = BORDER
    v.freeze_panes = "A2"

    # ---------------------------------------------------- 4. 統計摘要（全部用公式）
    s = wb.create_sheet("統計摘要")
    M = "需求單主檔"
    ai_col = get_column_letter(cols.index("ai_automation_level") + 1)
    eng_col = get_column_letter(cols.index("eng_effort_hours") + 1)
    prev_col = get_column_letter(cols.index("ai_preventable") + 1)
    res_col = get_column_letter(cols.index("resolution_type") + 1)
    rng = lambda col: f"'{M}'!${col}$2:${col}${n_rows}"  # noqa: E731

    s["A1"] = "統計摘要"
    s["A1"].font = TITLE_FONT
    s["A2"] = "所有數字皆為公式計算，修改主檔資料後會自動重算。"
    s["A2"].font = NOTE_FONT

    r = 4
    s.cell(row=r, column=1, value="總覽").font = Font(name=FONT, size=11, bold=True)
    r += 1
    overview = [
        ("需求單總數", f"=COUNTA({rng('A')})"),
        ("工程投入總工時", f"=ROUND(SUM({rng(eng_col)}),1)"),
        ("平均完整度分數", f"=ROUND(AVERAGE({rng(get_column_letter(comp_idx))}),1)"),
        ("平均釐清次數", f"=ROUND(AVERAGE({rng(clar_col)}),2)"),
        ("一次到位比例", f"=ROUND(COUNTIF({rng(clar_col)},0)/COUNTA({rng('A')}),3)"),
        ("釐清 3 次以上比例", f'=ROUND(COUNTIF({rng(clar_col)},">=3")/COUNTA({rng("A")}),3)'),
    ]
    for label, formula in overview:
        s.cell(row=r, column=1, value=label).font = BODY_FONT
        s.cell(row=r, column=2, value=formula).font = BODY_FONT
        if "比例" in label:
            s.cell(row=r, column=2).number_format = "0.0%"
        r += 1

    r += 1
    s.cell(row=r, column=1, value="AI 自動化分級分布").font = Font(name=FONT, size=11, bold=True)
    r += 1
    for h, c in zip(["分級", "工單數", "工單佔比", "工程工時", "工時佔比"], range(1, 6)):
        s.cell(row=r, column=c, value=h).font = Font(name=FONT, size=10, bold=True)
    r += 1
    for lvl in sorted(df["ai_automation_level"].unique()):
        s.cell(row=r, column=1, value=lvl).font = BODY_FONT
        s.cell(row=r, column=2, value=f"=COUNTIF({rng(ai_col)},A{r})").font = BODY_FONT
        s.cell(row=r, column=3, value=f"=B{r}/COUNTA({rng('A')})").font = BODY_FONT
        s.cell(row=r, column=3).number_format = "0.0%"
        s.cell(
            row=r, column=4, value=f"=ROUND(SUMIF({rng(ai_col)},A{r},{rng(eng_col)}),1)"
        ).font = BODY_FONT
        s.cell(row=r, column=5, value=f"=D{r}/SUM({rng(eng_col)})").font = BODY_FONT
        s.cell(row=r, column=5).number_format = "0.0%"
        r += 1

    r += 1
    s.cell(row=r, column=1, value="可預防性分布").font = Font(name=FONT, size=11, bold=True)
    r += 1
    for lvl in ["是", "部分", "否"]:
        s.cell(row=r, column=1, value=lvl).font = BODY_FONT
        s.cell(row=r, column=2, value=f"=COUNTIF({rng(prev_col)},A{r})").font = BODY_FONT
        s.cell(row=r, column=3, value=f"=B{r}/COUNTA({rng('A')})").font = BODY_FONT
        s.cell(row=r, column=3).number_format = "0.0%"
        r += 1

    r += 1
    s.cell(row=r, column=1, value="處理方式分布").font = Font(name=FONT, size=11, bold=True)
    r += 1
    for rt in sorted(df["resolution_type"].unique()):
        s.cell(row=r, column=1, value=rt).font = BODY_FONT
        s.cell(row=r, column=2, value=f"=COUNTIF({rng(res_col)},A{r})").font = BODY_FONT
        s.cell(row=r, column=3, value=f"=B{r}/COUNTA({rng('A')})").font = BODY_FONT
        s.cell(row=r, column=3).number_format = "0.0%"
        r += 1

    r += 1
    s.cell(
        row=r,
        column=1,
        value="資料為模擬生成，不含真實企業資料。生成器：src/adops_triage/generate_dataset.py",
    ).font = NOTE_FONT
    for w, c in zip([34, 14, 12, 14, 12], "ABCDE"):
        s.column_dimensions[c].width = w

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"✔ {OUT}（4 個工作表，{n_rows - 1} 筆）")


if __name__ == "__main__":
    main()
